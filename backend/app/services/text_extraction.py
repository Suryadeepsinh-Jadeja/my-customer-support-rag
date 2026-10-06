"""Turn an uploaded file into text, page by page.

PDFs: the embedded text layer per page; pages with (almost) no text are rendered and
OCR'd, so scanned PDFs work. Images: OCR. DOCX: paragraphs and tables. TXT: decoded.
"""

import asyncio
import io
import re
from dataclasses import dataclass

from app.services import file_validation as ft
from app.services.ocr import OcrEngine, OcrUnavailableError

# A page with fewer meaningful characters than this is treated as scanned.
MIN_TEXT_CHARS = 25
OCR_DPI = 200


class ExtractionError(RuntimeError):
    def __init__(self, code: str):
        super().__init__(code)
        self.code = code


@dataclass
class PageText:
    page_number: int
    text: str
    ocr: bool = False


def clean_text(text: str) -> str:
    text = text.replace("\x00", "").replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[ \t ]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return "\n".join(line.strip() for line in text.split("\n")).strip()


def _meaningful(text: str) -> int:
    return len(re.sub(r"\W", "", text))


def _pdf_pages(data: bytes, max_pages: int) -> list[tuple[int, str, bytes | None]]:
    """(page number, text layer, PNG render if the page needs OCR)."""
    import pymupdf

    try:
        doc = pymupdf.open(stream=data, filetype="pdf")
    except Exception as exc:  # PyMuPDF raises its own exception types
        raise ExtractionError("invalid_pdf") from exc
    with doc:
        if doc.needs_pass:
            raise ExtractionError("encrypted_pdf")
        if doc.page_count > max_pages:
            raise ExtractionError("too_many_pages")
        pages = []
        for index in range(1, doc.page_count + 1):
            page = doc[index - 1]
            text = page.get_text("text")
            render = None
            if _meaningful(text) < MIN_TEXT_CHARS:
                render = page.get_pixmap(dpi=OCR_DPI).tobytes("png")
            pages.append((index, text, render))
        return pages


def _docx_text(data: bytes) -> str:
    import docx

    try:
        document = docx.Document(io.BytesIO(data))
    except Exception as exc:
        raise ExtractionError("invalid_docx") from exc
    lines = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            lines.append(" | ".join(cell.text.strip() for cell in row.cells))
    return "\n".join(lines)


async def extract_pages(data: bytes, content_type: str, ocr: OcrEngine | None,
                        max_pages: int) -> list[PageText]:
    """Raises ExtractionError with a code the UI can explain."""
    if content_type == ft.TXT:
        return [PageText(1, clean_text(data.decode("utf-8")))]

    if content_type == ft.DOCX:
        return [PageText(1, clean_text(await asyncio.to_thread(_docx_text, data)))]

    if content_type in (ft.PNG, ft.JPEG):
        if ocr is None:
            raise ExtractionError("ocr_unavailable")
        try:
            return [PageText(1, clean_text(await ocr.image_to_text(data, content_type)), True)]
        except OcrUnavailableError as exc:
            raise ExtractionError("ocr_failed") from exc

    if content_type == ft.PDF:
        results = []
        for number, text, render in await asyncio.to_thread(_pdf_pages, data, max_pages):
            if render is not None and ocr is not None:
                try:
                    text = await ocr.image_to_text(render, ft.PNG)
                    results.append(PageText(number, clean_text(text), True))
                    continue
                except OcrUnavailableError as exc:
                    raise ExtractionError("ocr_failed") from exc
            results.append(PageText(number, clean_text(text)))
        if all(_meaningful(p.text) == 0 for p in results):
            raise ExtractionError("ocr_unavailable" if ocr is None else "no_text_found")
        return results

    raise ExtractionError("unsupported_file")
